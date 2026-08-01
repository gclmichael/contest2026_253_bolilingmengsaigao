from pathlib import Path
from datetime import date

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "submission" / "博丽灵梦赛高-FocusLoop-contest2026_253_bolilingmengsaigao"
ASSET_DIR = ROOT / "docs" / "report" / "assets"
DOCX_PATH = ROOT / "docs" / "report" / "FocusLoop_技术报告_可编辑备份.docx"

INK = "101311"
TEXT = "202621"
MUTED = "667068"
LIME = "D8FF58"
LIME_DARK = "35580A"
AMBER = "F6B84A"
PAPER = "F4F6F1"
LINE = "CED5CE"
WHITE = "FFFFFF"


def font(size, bold=False, color=TEXT):
    return {"size": Pt(size), "bold": bold, "color": RGBColor.from_string(color)}


def set_run(run, size=10, bold=False, color=TEXT, name="Microsoft YaHei"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def border(cell, color=LINE, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def cell_margin(cell, top=120, start=150, bottom=120, end=150):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn("w:" + name))
        if node is None:
            node = OxmlElement("w:" + name)
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("FOCUSLOOP  ·  ")
    set_run(run, 8, True, MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def add_paragraph(doc, text="", size=10, bold=False, color=TEXT, before=0, after=5, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_run(run, size, bold, color)
    return p


def add_rich_paragraph(doc, parts, size=10, before=0, after=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    for text, bold, color in parts:
        run = p.add_run(text)
        set_run(run, size, bold, color)
    return p


def add_heading(doc, number, title, kicker=None):
    if kicker:
        add_paragraph(doc, kicker.upper(), 8, True, LIME_DARK, after=2)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"{number}  ")
    set_run(r, 20, True, LIME_DARK)
    r = p.add_run(title)
    set_run(r, 20, True, INK)
    return p


def add_subheading(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title)
    set_run(r, 11.5, True, INK)
    return p


def add_bullets(doc, items, size=9.3):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.45)
        p.paragraph_format.first_line_indent = Cm(-0.2)
        p.paragraph_format.space_after = Pt(2.5)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(item)
        set_run(r, size, False, TEXT)


def add_band(doc, title, text, accent=LIME, dark=False):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(0.4)
    table.columns[1].width = Cm(16.8)
    left, body = table.rows[0].cells
    shade(left, accent)
    shade(body, INK if dark else PAPER)
    for c in (left, body):
        border(c, INK if dark else PAPER, "0")
        cell_margin(c, 120, 130, 120, 130)
    p = body.paragraphs[0]
    r = p.add_run(title)
    set_run(r, 10.5, True, LIME if dark else INK)
    r = p.add_run("\n" + text)
    set_run(r, 9, False, WHITE if dark else TEXT)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_table(doc, headers, rows, widths=None, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    if widths:
        for idx, width in enumerate(widths):
            table.columns[idx].width = Cm(width)
    head = table.rows[0]
    set_repeat_table_header(head)
    for idx, text in enumerate(headers):
        cell = head.cells[idx]
        shade(cell, INK)
        border(cell, INK)
        cell_margin(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        r = p.add_run(text)
        set_run(r, font_size, True, LIME)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            cell = cells[idx]
            shade(cell, WHITE if ridx % 2 == 0 else PAPER)
            border(cell)
            cell_margin(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(text))
            set_run(r, font_size, idx == 0, TEXT)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_metrics(doc, metrics):
    table = doc.add_table(rows=1, cols=len(metrics))
    table.autofit = False
    for idx, (value, label) in enumerate(metrics):
        cell = table.cell(0, idx)
        shade(cell, INK)
        border(cell, WHITE, "0")
        cell_margin(cell, 170, 140, 170, 140)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(value)
        set_run(r, 18, True, LIME)
        r = p.add_run("\n" + label)
        set_run(r, 8, False, WHITE)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_image(doc, path, width_cm, caption=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    if caption:
        cap = add_paragraph(doc, caption, 7.5, False, MUTED, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
        cap.paragraph_format.keep_with_next = False


def add_two_images(doc, left, right, left_caption, right_caption, width_cm=7.8):
    table = doc.add_table(rows=2, cols=2)
    table.autofit = False
    for cell in table.rows[0].cells + table.rows[1].cells:
        border(cell, WHITE, "0")
        cell_margin(cell, 50, 80, 50, 80)
    for idx, path in enumerate((left, right)):
        p = table.cell(0, idx).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Cm(width_cm))
    for idx, text in enumerate((left_caption, right_caption)):
        p = table.cell(1, idx).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_run(r, 7.5, False, MUTED)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_page_break(doc):
    doc.add_page_break()


def make_flow(path, nodes, subtitle, width=1800, height=600):
    font_path = Path("C:/Windows/Fonts/msyh.ttc")
    title_font = ImageFont.truetype(str(font_path), 54)
    node_font = ImageFont.truetype(str(font_path), 31)
    small_font = ImageFont.truetype(str(font_path), 24)
    image = Image.new("RGB", (width, height), "#F4F6F1")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, width, 18), fill="#D8FF58")
    draw.text((55, 47), subtitle, font=title_font, fill="#101311")
    count = len(nodes)
    gap = 35
    x0 = 55
    y0 = 180
    box_w = int((width - 110 - gap * (count - 1)) / count)
    box_h = 285
    for idx, (title, body) in enumerate(nodes):
        x = x0 + idx * (box_w + gap)
        fill = "#101311" if idx % 2 == 0 else "#FFFFFF"
        text_color = "#FFFFFF" if idx % 2 == 0 else "#202621"
        border_color = "#101311"
        draw.rectangle((x, y0, x + box_w, y0 + box_h), fill=fill, outline=border_color, width=4)
        draw.text((x + 24, y0 + 26), title, font=node_font, fill="#D8FF58" if idx % 2 == 0 else "#35580A")
        lines = body.split("\n")
        for line_idx, line in enumerate(lines):
            draw.text((x + 24, y0 + 92 + line_idx * 42), line, font=small_font, fill=text_color)
        if idx < count - 1:
            ax = x + box_w + 7
            ay = y0 + box_h // 2
            draw.line((ax, ay, ax + gap - 14, ay), fill="#667068", width=6)
            draw.polygon(((ax + gap - 14, ay - 11), (ax + gap - 3, ay), (ax + gap - 14, ay + 11)), fill="#667068")
    image.save(path)


def prepare_assets():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    make_flow(
        ASSET_DIR / "architecture.png",
        [
            ("输入", "语音主题\n本地材料"),
            ("模型", "知识提取\n题卡生成"),
            ("本地", "校验判分\n进度调度"),
            ("情境", "压力 + 活动\n反馈自适应"),
            ("执行", "cron_add\nlaunch_quickapp"),
        ],
        "五层协同架构",
    )
    make_flow(
        ASSET_DIR / "grounding.png",
        [
            ("材料", "限定本地收件箱\nUTF-8 文本"),
            ("提取", "知识点 + 原文依据\n3–8 张题卡"),
            ("校验", "字段 / 长度\n选项 / 答案索引"),
            ("呈现", "问题 + 解释\n依据可核对"),
        ],
        "材料到题卡：每一步都可检查",
    )
    make_flow(
        ASSET_DIR / "context.png",
        [
            ("记忆", "到期或 30 分钟内\n已有学习记录"),
            ("情境", "压力低\n连续静止"),
            ("约束", "用户已开启\n冷却期结束"),
            ("反馈", "开始复习 / 稍后\n调整个人阈值"),
        ],
        "从固定提醒到个人化短时复习窗口",
    )


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.35)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(1.55)
    section.right_margin = Cm(1.55)
    section.header_distance = Cm(0.45)
    section.footer_distance = Cm(0.55)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(TEXT)

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    header = section.header
    p = header.paragraphs[0]
    p.text = "2026 首届 openvela AI 硬件开发者大赛  ·  技术报告"
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in p.runs:
        set_run(run, 7.5, True, MUTED)
    page_number(section.footer.paragraphs[0])


def build_report():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prepare_assets()
    doc = Document()
    configure_document(doc)

    # 1. Cover
    cover = doc.add_table(rows=1, cols=1)
    cover.autofit = False
    cell = cover.cell(0, 0)
    shade(cell, INK)
    border(cell, INK, "0")
    cell_margin(cell, 650, 650, 650, 650)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.add_run().add_picture(str(ROOT / "quickapp/focusloop/src/common/logo.png"), width=Cm(2.4))
    r = p.add_run("\n\nFOCUSLOOP")
    set_run(r, 34, True, WHITE)
    r = p.add_run("\n腕上主动学习闭环")
    set_run(r, 22, True, LIME)
    r = p.add_run("\n\n把用户材料整理成可核对的短题，\n在更合适的时机完成下一次复习。")
    set_run(r, 13, False, WHITE)
    r = p.add_run("\n\n\n2026 首届 openvela AI 硬件开发者大赛\n队伍：博丽灵梦赛高  ·  编号：253\n赛道：手表应用创新 + AI 硬件产品创新")
    set_run(r, 9.5, False, "B9C2B9")
    for _ in range(2):
        cell.add_paragraph()
    shot_p = cell.add_paragraph()
    shot_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shot_p.add_run().add_picture(str(ROOT / "docs/screenshots/focusloop-goldfish-dashboard.png"), width=Cm(8.3))
    add_page_break(doc)

    # 2. Submission info and abstract
    add_heading(doc, "01", "作品信息与摘要", "作品概览")
    add_table(doc, ["字段", "内容"], [
        ("作品名称", "FocusLoop 腕上主动学习闭环"),
        ("队伍名称", "博丽灵梦赛高（253）"),
        ("成员与分工", "龚城立（队长）：产品定义、系统集成、测试与展示"),
        ("参赛方向", "手表应用创新 + AI 硬件产品创新"),
        ("开源许可", "Apache License 2.0"),
    ], [3.4, 13.6], 9)
    add_subheading(doc, "作品摘要（300 字以内）")
    abstract = (
        "FocusLoop 面向学习计划难以持续、固定提醒容易打扰的问题，在 openvela 手表上打通语音或材料输入、"
        "知识题卡生成、专注、主动回忆和再次触达。大模型提取知识点并生成带原文依据的短题；本地代码负责结构校验、"
        "判分和复习调度；用户开启智能窗口后，手表结合压力、活动状态、到期时间和反馈记录选择短时复习时机。"
        "当前版本通过 26 项测试、20 轮连续测试和 5 轮连续构建，生产依赖漏洞为 0，可在模型或健康服务不可用时保持核心流程。"
    )
    assert len(abstract) <= 300
    add_band(doc, "摘要", abstract, LIME, True)
    add_subheading(doc, "评委快速定位")
    add_table(doc, ["评分维度", "报告证据", "对应页"], [
        ("技术难度 30", "五层架构、Agent 协议、状态恢复、openvela 接口", "4、5、8"),
        ("创新性 20", "材料依据题卡 + 情境推荐 + 用户反馈自适应", "3、5、6"),
        ("完成度 20", "完整 UI 流程、RPK、连续测试与构建", "7、9"),
        ("AI 开发 10", "运行时 Skill、开发验证 Skill、AI-Native 说明", "8"),
        ("商业价值 10", "个人学习与机构微课的落地路径", "10"),
        ("展示效果 10", "Goldfish 运行画面、演示视频、可复现证据", "7、9"),
    ], [3.2, 10.6, 2.2], 8.2)
    add_page_break(doc)

    # 3. Problem and innovation
    add_heading(doc, "02", "问题、挑战与核心创新", "问题与创新")
    add_rich_paragraph(doc, [
        ("现实痛点：", True, INK),
        ("内容生成后缺少后续执行；固定提醒不理解用户是否方便；长材料难以适配几十秒的腕上交互。", False, TEXT),
    ], 11, after=7)
    add_metrics(doc, [("5 步", "输入到再次触达"), ("30 秒", "可完成的短时回忆"), ("默认关闭", "情境推荐由用户选择"), ("本地", "健康数据只做时机判断")])
    add_image(doc, ROOT / "docs/screenshots/focusloop-overview.png", 16.6, "真实产品流程：计划、专注、答题、再次触达")
    add_subheading(doc, "三项有明确增量的创新")
    add_table(doc, ["创新", "常见方案", "FocusLoop 的增量"], [
        ("可核对题卡", "模型生成题目，来源不透明", "每张导入题卡保留材料依据；界面同时展示解释与依据"),
        ("情境复习窗口", "只按遗忘间隔或固定时间提醒", "同时检查到期、压力、活动、冷却期，条件不合适时保持安静"),
        ("反馈自适应", "统一压力阈值", "用户选择“开始复习”或“稍后”，本地更新个人阈值与接受率"),
    ], [3.1, 5.3, 8.1], 8.2)
    add_page_break(doc)

    # 4. Architecture
    add_heading(doc, "03", "总体架构与运行流程", "系统方案")
    add_image(doc, ASSET_DIR / "architecture.png", 17.1)
    add_table(doc, ["层", "关键模块", "职责与降级"], [
        ("输入层", "PTT/ASR、本地材料收件箱", "接收主题或材料；语音不可用时仍可选预设主题"),
        ("模型层", "ai_agent + FocusLoop Skill", "知识提取、题卡生成、定时工具调用；超时切换离线计划"),
        ("本地层", "agent_protocol / scheduler / store", "校验、判分、到期计算、状态恢复，完全可测试"),
        ("情境层", "service.health / system.sensor", "只在用户开启后判断复习时机；能力缺失不影响核心流程"),
        ("执行层", "cron_add / launch_quickapp", "持久定时并主动拉起 QuickApp"),
    ], [2.3, 5.1, 9.2], 8.1)
    add_band(doc, "状态恢复", "下一次复习先写入本地 pendingScheduleAt。Agent 调度失败或应用重启后，首页继续同步；只有收到结构化 scheduled:true 才清除待同步状态。", AMBER)
    add_subheading(doc, "一次学习的状态流")
    add_paragraph(doc, "无计划 → 生成计划 → 单任务专注 → 主动回忆 → 更新间隔 → 持久定时 → 到期拉起 → 再次回忆", 10.5, True, LIME_DARK, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_page_break(doc)

    # 5. Model and grounding
    add_heading(doc, "04", "大模型链路：从材料理解到可执行题卡", "AI 算法")
    add_image(doc, ASSET_DIR / "grounding.png", 17.1)
    add_table(doc, ["环节", "输入 / 输出", "约束"], [
        ("CREATE_PLAN", "主题 → 学习目标 + 3–8 张题卡", "单轮快速文本模式；输出放在 focusloop-json 标签中"),
        ("IMPORT", "固定收件箱材料 → 知识点 + 题卡 + 原文依据", "只读指定路径；材料不能改变工具权限或操作类型"),
        ("本地规范化", "问题、3 个选项、答案、解释、概念、依据", "长度、数量、答案索引和字段类型全部检查"),
        ("异常恢复", "超时、缺字段、服务未配置", "切换内置题目，专注、答题和进度保存继续可用"),
    ], [3, 6.5, 7.1], 8.2)
    add_subheading(doc, "为什么需要模型")
    add_bullets(doc, [
        "同一份课程笔记可能包含定义、比较、流程和例外条件，模型负责把非结构化内容拆成适合回忆的知识单元。",
        "生成结果带解释和材料依据，评委与用户可以检查答案来自哪里，降低题卡看似合理但缺乏支撑的风险。",
        "定时、判分和复习间隔不交给模型；这些状态变化由本地确定性代码完成，便于测试与恢复。",
    ], 9.5)
    add_band(doc, "模型配置", "OpenAI 兼容接口；当前测试使用 MiMo-V2.5。密钥由设备环境配置，不写入仓库、报告或构建产物。", LIME, True)
    add_page_break(doc)

    # 6. Context algorithm
    add_heading(doc, "05", "情境算法：找到更合适的短时复习窗口", "情境判断")
    add_image(doc, ASSET_DIR / "context.png", 17.1)
    add_subheading(doc, "本地判定条件")
    add_table(doc, ["信号", "规则", "目的"], [
        ("用户选择", "智能推荐默认关闭，开启后才订阅数据", "把控制权交给用户"),
        ("记忆状态", "至少完成一次学习；题卡已到期或 30 分钟内到期", "只推荐真实需要复习的内容"),
        ("压力", "低于个人阈值，初始值 25，限制在 15–35", "避开压力较高时刻"),
        ("活动", "至少 6 个加速度采样显示连续静止", "避开走路或运动"),
        ("节流", "每次推荐后进入 30 分钟冷却", "减少连续打扰"),
        ("反馈", "开始复习 / 稍后，以 80/20 平滑更新阈值", "逐步适应个人接受习惯"),
    ], [2.4, 8.1, 6.1], 8.1)
    add_two_images(
        doc,
        ROOT / "docs/screenshots/focusloop-smart-overview.png",
        ROOT / "docs/screenshots/focusloop-goldfish-quiz-result.png",
        "智能窗口：关闭、等待与可复习状态",
        "答题反馈：解释与复习间隔更新",
        7.4,
    )
    add_band(doc, "隐私边界", "压力与加速度只在本地判断时机，不进入模型，也不用于医疗结论。后台驻留时监测持续；进程结束后由持久定时任务保障时间型提醒。", AMBER)
    add_page_break(doc)

    # 7. UI and complete flow
    add_heading(doc, "06", "软件实现与腕上交互", "产品实现")
    add_paragraph(doc, "界面针对 466×466 圆形表盘设计。每页保留一个主动作，状态文字短而明确，常用流程可在数次点击内完成。", 10.5, False, TEXT, after=6)
    add_two_images(
        doc,
        ROOT / "docs/screenshots/focusloop-goldfish-plan.png",
        ROOT / "docs/screenshots/focusloop-goldfish-focus.png",
        "计划页：语音、预设主题、材料导入",
        "专注页：15 / 25 / 40 分钟单任务计时",
        7.6,
    )
    add_two_images(
        doc,
        ROOT / "docs/screenshots/focusloop-goldfish-quiz.png",
        ROOT / "docs/screenshots/focusloop-goldfish-dashboard.png",
        "答题页：三选一主动回忆",
        "首页：进度、下一次到期和立即复习",
        7.6,
    )
    add_table(doc, ["页面", "主要动作", "异常状态"], [
        ("计划", "语音主题、预设主题、导入材料", "ASR 或模型不可用时给出状态并启用离线题目"),
        ("专注", "开始 / 暂停 / 完成", "振动能力缺失不阻断计时"),
        ("答题", "选择答案、查看解释和依据", "进度先本地保存，再同步 Agent 定时"),
        ("智能窗口", "开启、开始复习、稍后", "健康能力缺失时单页降级"),
    ], [2.5, 6.2, 7.9], 8.1)
    add_page_break(doc)

    # 8. openvela, Skills and AI Native
    add_heading(doc, "07", "openvela 深度集成与开发方法", "平台集成")
    add_table(doc, ["openvela 能力", "使用方式", "产品价值"], [
        ("system.velaclaw", "QuickApp 向 ai_agent 发送计划、语音、调度请求", "把腕上交互接入 Agent"),
        ("ai_agent Skill", "约束 SCHEDULE / REVIEW / IMPORT 的工具与文件范围", "让模型输出转成受控系统动作"),
        ("cron_add", "保存一次性复习任务并用 cron_list 去重", "应用退出后仍能按时执行"),
        ("launch_quickapp", "到期主动打开 FocusLoop", "减少用户重新寻找入口的成本"),
        ("service.health", "读取实时压力样本", "为低打扰时机提供设备信号"),
        ("system.sensor", "连续读取加速度并判断静止", "避开运动中的提醒"),
        ("system.storage", "保存计划、卡片间隔、待同步任务和偏好", "支持重启恢复"),
    ], [3.4, 7.3, 5.9], 8.05)
    add_subheading(doc, "两个可复用 Skill")
    add_table(doc, ["Skill", "用途", "复用价值"], [
        ("FocusLoop runtime Skill", "规定导入、复习与调度操作；验证真实工具结果", "其他 Agent 应用可复用受控调度和本地收件箱模式"),
        ("focusloop-verify", "测试、构建、漏洞、RPK、敏感信息和文案检查", "评审前一条命令完成交付核对，不会上传或改写历史"),
    ], [4, 7, 5.6], 8.2)
    add_subheading(doc, "AI-Native 开发说明")
    add_table(doc, ["模板字段", "填写内容"], [
        ("AI Coding 代码占比", "约 50%，以当前参赛分支新增/修改代码为口径估算；关键架构、设备联调、验收和提交内容均由参赛者复核"),
        ("AI 编程工具", "Codex：代码实现、测试、文档整理与交付检查"),
        ("MCP 使用", "未使用 VelaJS MCP；通过本地 Shell、ADB 与官方文档完成联调"),
        ("Skills 使用", "运行时 FocusLoop Skill；开发侧 focusloop-verify Skill"),
        ("Token 使用量", "MiMo：2,193,625,712+ Token；GPT：辅助使用，未统计"),
    ], [4.3, 12.3], 8.2)
    add_page_break(doc)

    # 9. Quantified testing
    add_heading(doc, "08", "量化测试、性能与可靠性", "测试验证")
    add_metrics(doc, [("26/26", "单轮测试"), ("20/20", "连续测试"), ("5/5", "连续构建"), ("0", "生产依赖漏洞")])
    add_table(doc, ["测试类别", "环境与方法", "结果"], [
        ("功能测试", "Node.js 测试覆盖协议解析、提示词边界、复习算法、导入、语音、情境判断和页面约束", "26 项全部通过"),
        ("重复稳定性", "Windows 11 / Node 24.14.1；连续运行 20 轮测试", "共 520 项检查，0 项失败"),
        ("构建稳定性", "AIoT Toolkit 2.0.5；连续运行 5 轮 debug 构建", "5/5 成功"),
        ("制品", "debug RPK + SHA-256", "约 72.4 KB；制品与哈希位于仓库 artifacts/"),
        ("依赖安全", "npm audit --omit=dev", "0 个生产依赖漏洞"),
        ("Goldfish", "Ubuntu 22.04，xiaomi_watch_s1 / arm64 Goldfish", "系统 3561 个目标构建完成；核心页面与离线链路已运行"),
    ], [3.2, 9.2, 4.2], 8.1)
    add_subheading(doc, "故障隔离矩阵")
    add_table(doc, ["能力不可用", "受影响部分", "仍可用部分"], [
        ("模型 / 网络", "新主题的在线生成", "内置题目、专注、答题、进度保存"),
        ("语音 / ASR", "语音主题输入", "预设主题和材料导入"),
        ("service.health", "情境推荐", "时间型复习与完整学习流程"),
        ("Agent 调度暂时失败", "主动拉起等待重试", "pendingScheduleAt 保留，首页恢复同步"),
    ], [3.6, 5.2, 7.8], 8.1)
    add_band(doc, "尚待设备侧补测", "真实手表的内存、功耗、ASR 端到端延迟，以及带健康服务镜像的压力订阅运行数据。报告不把编译或 Mock 数据写成真机结论。", AMBER)
    add_page_break(doc)

    # 10. User value and validation
    add_heading(doc, "09", "目标用户、验证计划与当前边界", "应用价值")
    add_table(doc, ["首批场景", "目标用户", "可验证价值"], [
        ("语言与认证备考", "需要高频短时回忆的学生与职场人", "复习完成率、次日留存、7 日留存"),
        ("技术知识复习", "开发者、工程师、实验室成员", "把个人文档快速转成可复习题卡"),
        ("企业微课", "需要岗位知识与安全规范训练的机构", "课程材料统一下发，腕上完成低打扰训练"),
    ], [3.5, 6.5, 6.6], 8.3)
    add_subheading(doc, "两周验证路线")
    add_table(doc, ["阶段", "工作", "指标"], [
        ("两周试用", "20–50 名校内用户完成真实学习任务", "建议接受率、复习完成率、次日与 7 日留存"),
        ("设备验收", "miwear 健康镜像与兼容手表联调", "压力订阅稳定性、误打扰率、功耗"),
        ("内容质量", "抽检材料生成的题卡、答案和依据", "材料转卡成功率、题卡人工通过率"),
        ("继续投入条件", "根据试用和真机数据决定是否扩展机构课程", "接受率、完成率与留存达到预设目标"),
    ], [2.8, 8.5, 5.3], 8.2)
    add_subheading(doc, "当前限制")
    add_bullets(doc, [
        "材料导入目前通过 ADB 推送到固定收件箱，面向普通用户的文件选择器需要后续设备 API 支持。",
        "情境模型当前使用轻量规则与反馈调整，尚未用真实用户数据训练；先保证可解释、可关闭和不打扰。",
        "后台情境监测依赖应用进程驻留；进程结束后只保留持久的时间型提醒。",
    ], 9.2)
    add_band(doc, "落地判断", "先用真实学习任务验证题卡质量、建议接受率、复习完成率和留存。数据成立后，再扩展课程管理和机构部署，避免在缺少用户证据时提前堆叠商业功能。", LIME, True)

    props = doc.core_properties
    props.title = "FocusLoop 腕上主动学习闭环 - 技术报告"
    props.subject = "2026 首届 openvela AI 硬件开发者大赛"
    props.author = "博丽灵梦赛高"
    props.keywords = "openvela, FocusLoop, AI Agent, wearable, spaced repetition"
    props.comments = "队伍 253"
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build_report()
